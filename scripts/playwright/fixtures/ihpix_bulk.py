"""Preparar solo la base local desechable ihpix_bulk para la prueba del navegador."""

import io
import json
from pathlib import Path
from openpyxl import load_workbook
from sqlalchemy import text
from ckan.cli import load_config
from ckan.config.middleware import make_app

app = make_app(load_config("/qa/app.ini"))
import ckan.model as model
import ckan.plugins.toolkit as tk
from ckanext.theme_ejemplo import (
    ihpix_bulk as B,
    ihpix_forms as F,
    ihpix_bulk_schema as S,
)

assert tk.config["sqlalchemy.url"].endswith("@ihpix-bulk-db/ihpix_bulk")
model.Session.execute(text("TRUNCATE TABLE ihpix_import_batch, ihpix_activity CASCADE"))
model.Session.commit()
root = Path("/qa/spreadsheets")
root.mkdir(exist_ok=True)
for locale in ("en", "es", "fr", "ar"):
    raw = B.template(locale=locale, example=True)
    (root / (locale + "-example.xlsx")).write_bytes(raw)
    (root / (locale + "-template.xlsx")).write_bytes(B.template(locale=locale))
    book = load_workbook(io.BytesIO(raw))
    name = next(
        row[1].value
        for row in book["_IHP_META"].iter_rows(min_row=5, max_row=14)
        if row[0].value == "Activities"
    )
    book[name]["A5"] = book[name]["A5"].value + " (preview)"
    output = io.BytesIO()
    book.save(output)
    (root / (locale + "-review.xlsx")).write_bytes(output.getvalue())
book = load_workbook(root / "en-example.xlsx")
fields = [cell.value for cell in book["Activities"][1]]
invalid_cell = book["Activities"].cell(5, fields.index("num_knowledge_products") + 1)
invalid_cell.value = -1
(root / "invalid-cell.txt").write_text(invalid_cell.coordinate)
output = io.BytesIO()
book.save(output)
(root / "invalid.xlsx").write_bytes(output.getvalue())
book = load_workbook(io.BytesIO(B.template()))
book["Activities"]["A5"] = "Minimal imported draft"
output = io.BytesIO()
book.save(output)
(root / "minimal.xlsx").write_bytes(output.getvalue())
book = load_workbook(io.BytesIO(B.template_v1()))
book["Activities"].append(["Legacy1", "Legacy draft"])
output = io.BytesIO()
book.save(output)
(root / "legacy.xlsx").write_bytes(output.getvalue())
book = load_workbook(io.BytesIO(B.template_v2()))
book["Activities"]["B5"] = "Draft from the previous guided template"
output = io.BytesIO()
book.save(output)
(root / "legacy-v2.xlsx").write_bytes(output.getvalue())
expected = [
    F.validate_report_payload(row["data"], True)
    for row in B.parse((root / "en-example.xlsx").read_bytes())["rows"]
]
(root / "expected.json").write_text(
    json.dumps(expected, default=lambda value: value.isoformat())
)
print("Isolated fixtures ready")

book = load_workbook(io.BytesIO(B.template()))
for row, output in [(5, "1.1"), (6, "1.10")]:
    fields = [cell.value for cell in book["Activities"][1]]
    for field, value in {
        "title": "Rollback activity " + output,
        "priority_area": "PA1",
        "output": output,
    }.items():
        book["Activities"].cell(row, fields.index(field) + 1, value)
raw = io.BytesIO()
book.save(raw)
(root / "rollback.xlsx").write_bytes(raw.getvalue())
(root / "ui-labels.json").write_text(
    json.dumps(
        {
            locale: S.translator(locale)("Import activities from Excel")
            for locale in ("en", "es", "fr", "ar")
        }
    )
)
