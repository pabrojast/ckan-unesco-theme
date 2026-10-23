# -*- coding: utf-8 -*-
"""Plantilla XLSX y validación sin acceso a CKAN ni escrituras de datos."""
import datetime
import io
import json
import re
import zipfile
from xml.etree.ElementTree import ParseError

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from ckanext.theme_ejemplo import ihpix_forms as F, ihpix_constants as C
from ckanext.theme_ejemplo import ihpix_relations as R, ihpix_links as L

VERSION = 'IHP-IX/1'
MAX_ROWS = 2000
FIELDS = ['row_ref'] + list(dict.fromkeys(
    ['title', 'priority_area'] + list(F.TEXT_FIELDS) + list(F.INT_FIELDS) +
    list(F.BOOL_FIELDS) + list(F.FORM_DATE_FIELDS) + list(F.MULTI_FORM_FIELDS)))
SHEETS = {
    'Activities': FIELDS,
    'Contributions': ['row_ref', 'priority_area', 'output', 'key_activity_code'],
    'Synergies': ['row_ref', 'priority_area', 'output', 'key_activity_code'],
    'Supporters': ['row_ref', 'country', 'types'],
    'Partners': ['row_ref', 'organization_id'],
    'Resources': ['row_ref', 'link_type', 'target_kind', 'target_id', 'title', 'url', 'description', 'event_date'],
}


def template_v1(organizations=()):
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = 'Instructions'
    instructions.append(['Template version', VERSION])
    for text in (
        'One activity per row. row_ref must be unique (e.g. A001).',
        'Use row_ref on relation sheets to link multiple entries to an activity.',
        'Use yes/no for gates. Empty answers remain unanswered.',
        'Separate multiple choices or support types with semicolons.',
        'Use ISO dates YYYY-MM-DD and nonnegative whole numbers up to 2147483647.',
        'The first Contributions entry is the primary classification. Synergies do not increase contribution counts.',
        'Support types: financial;in_kind;hosting. Country codes are in Countries.',
        'Organization identifiers are in Organizations. External institutions and partners use free text.',
        'Resources can reference an existing item or an external HTTP(S) URL. Upload binary files after importing drafts.',
        'Upload, validate and preview before confirming. Import creates drafts only. Repeated confirmation is idempotent.',
        'Incomplete submission fields are warnings; malformed values prevent import.',
    ):
        instructions.append([text])
    for name, headers in SHEETS.items():
        workbook.create_sheet(name).append(headers)
    vocabulary = workbook.create_sheet('Catalogs')
    vocabulary.append(['field', 'value', 'label'])
    for field, choices in list(F.MULTI_FIELDS.items()) + [
        ('priority_area', C.PRIORITY_AREAS), ('biennium', C.BIENNIA),
        ('institution_type', C.LEAD_INSTITUTION_TYPES)]:
        for choice in choices or ():
            vocabulary.append([field, choice, C.flagship_label(choice) if field == 'flagships' else choice])
    taxonomy = workbook.create_sheet('Taxonomy')
    taxonomy.append(['priority_area', 'output', 'output_title', 'key_activity_code', 'key_activity_title'])
    for output, entries in C.load_key_activities().items():
        for item in entries:
            taxonomy.append([C.priority_area_for_output(output), output, C.output_title(output), item['code'], item['title']])
    countries = workbook.create_sheet('Countries')
    countries.append(['iso2', 'name', 'electoral_group', 'regions'])
    for country in R.countries():
        countries.append([country['iso2'], country['country_title_en'], country['electoral_group'], ';'.join(country['regional_group'] or [])])
    orgs = workbook.create_sheet('Organizations')
    orgs.append(['organization_id', 'title'])
    for org in organizations:
        orgs.append([org['id'], org.get('title') or org.get('name')])
    for sheet in workbook:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 28
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # Una etiqueta nunca debe ejecutarse como fórmula.
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def parse(raw, defaults=None):
    try:
        return _parse(raw, defaults)
    except (zipfile.BadZipFile, KeyError, OSError, ParseError, SyntaxError,
            RuntimeError, NotImplementedError) as exc:
        # Las hojas se leen de forma diferida; también pueden fallar al iterar.
        raise ValueError('Upload a valid XLSX workbook') from exc


def _parse(raw, defaults=None):
    errors, warnings, activities = [], [], {}
    if len(raw) > 5 * 1024 * 1024:
        raise ValueError('Workbook exceeds 5 MB')
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 30 * 1024 * 1024:
                raise ValueError('Expanded workbook exceeds 30 MB')
            # No confiar en <dimension>: algunos escritores omiten filas pobladas.
            # Acotar las coordenadas reales antes de que openpyxl construya filas.
            from openpyxl.xml.functions import iterparse
            from openpyxl.utils.cell import column_index_from_string
            for name in archive.namelist():
                if not re.fullmatch(r'xl/worksheets/sheet[0-9]+\.xml', name):
                    continue
                with archive.open(name) as stream:
                    previous_row, current_row, previous_column = 0, None, 0
                    for event, element in iterparse(stream, events=('start', 'end')):
                        tag = element.tag.rsplit('}', 1)[-1]
                        if event == 'start' and tag == 'row':
                            current_row = int(element.get('r', str(previous_row + 1)))
                            if current_row <= previous_row:
                                raise ValueError('Worksheet structure changed; download the current template')
                            if current_row > MAX_ROWS * 8 + 4:
                                raise ValueError('Worksheet dimensions exceed template limits')
                            previous_row, previous_column = current_row, 0
                        if event == 'start' and tag == 'c':
                            address = re.fullmatch(r'([A-Z]{1,3})([0-9]+)', element.get('r', ''))
                            if not address or column_index_from_string(address[1]) > len(FIELDS) or int(address[2]) > MAX_ROWS * 8 + 4:
                                raise ValueError('Worksheet dimensions exceed template limits')
                            column = column_index_from_string(address[1])
                            if int(address[2]) != current_row or column <= previous_column:
                                raise ValueError('Worksheet structure changed; download the current template')
                            previous_column = column
                        if event == 'end':
                            if tag == 'row':
                                current_row = None
                            element.clear()
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        for sheet in workbook:
            sheet.reset_dimensions()
            sheet.calculate_dimension(force=True)
    except (zipfile.BadZipFile, KeyError, OSError, ParseError, SyntaxError,
            RuntimeError, NotImplementedError) as exc:
        raise ValueError('Upload a valid XLSX workbook') from exc
    if '_IHP_META' in workbook:
        try:
            from ckanext.theme_ejemplo.ihpix_bulk_workbook import parse_v2
            if workbook["_IHP_META"]["B1"].value == "IHP-IX/3":
                from ckanext.theme_ejemplo import ihpix_bulk_flat
                return parse_v2(workbook, schema=ihpix_bulk_flat)
            return parse_v2(workbook)
        finally:
            workbook.close()
    if 'Instructions' not in workbook or workbook['Instructions']['B1'].value != VERSION:
        workbook.close()
        raise ValueError('Unsupported template version; download the current template')
    def error(sheet, row, field, message, target=errors):
        headers = SHEETS[sheet]
        from openpyxl.utils import get_column_letter
        cell = get_column_letter(headers.index(field) + 1) + str(row) if field in headers else str(row)
        target.append({'sheet': sheet, 'row': row, 'cell': cell, 'field': field, 'message': message})
    records = {}
    try:
        for name, headers in SHEETS.items():
            if name not in workbook:
                raise ValueError('Missing worksheet: ' + name)
            sheet = workbook[name]
            if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > len(headers):
                raise ValueError('Worksheet dimensions exceed template limits: ' + name)
            rows = sheet.iter_rows()
            first = next(rows, ())
            if [c.value for c in first] != headers:
                raise ValueError('Column headers do not match template: ' + name)
            records[name] = []
            for number, cells in enumerate(rows, 2):
                if not any(c.value is not None for c in cells):
                    continue
                values = {}
                for field, cell in zip(headers, cells):
                    if cell.data_type == 'f':
                        error(name, number, field, 'Formulas are not allowed; paste values')
                    value = cell.value
                    if isinstance(value, datetime.datetime):
                        value = value.date().isoformat()
                    elif isinstance(value, datetime.date):
                        value = value.isoformat()
                    elif isinstance(value, float) and value.is_integer():
                        value = int(value)
                    values[field] = '' if value is None else str(value).strip()
                records[name].append((number, values))
        return validate_records(records, defaults, initial_errors=errors)
    finally:
        workbook.close()

def validate_records(records, defaults=None, sources=None, version=VERSION, initial_errors=None):
    errors, warnings, activities = list(initial_errors or []), [], {}
    sources = dict(sources or {})
    original_records = {name: [(row, dict(data)) for row, data in entries] for name, entries in records.items()}
    def location(sheet, row, field):
        from openpyxl.utils import get_column_letter
        headers = SHEETS[sheet]
        fallback = {'sheet': sheet, 'row': row,
                    'cell': get_column_letter(headers.index(field) + 1) + str(row) if field in headers else str(row)}
        return (sources or {}).get((sheet, row, field), fallback)
    def error(sheet, row, field, message, target=errors):
        target.append(dict(location(sheet, row, field), field=field, message=message,
                           _sheet=sheet, _row=row))
    for number, data in records['Activities']:
        ref = data.pop('row_ref')
        if not ref or ref in activities:
            error('Activities', number, 'row_ref', 'Use a unique, nonempty row reference')
            continue
        for field, value in (defaults or {}).items():
            if field in ('institution', 'institution_org_id'):
                continue
            if not data.get(field):
                data[field] = value
        # No atribuir una institución externa al perfil por completar sólo el ID.
        if not data.get('institution') and not data.get('institution_org_id'):
            for field in ('institution', 'institution_org_id'):
                data[field] = (defaults or {}).get(field, '')
        for field in F.MULTI_FORM_FIELDS:
            data[field] = [v.strip() for v in data[field].split(';') if v.strip()]
        activities[ref] = {'row': number, 'data': data, 'resource_rows': [], 'partner_rows': [],
                           'sources': {field: location('Activities', number, field) for field in FIELDS}}
    invalid_relations = set()
    for name, field in [('Contributions','contributions_json'), ('Synergies','synergies_json'),
                        ('Supporters','supporters_json'), ('Partners','partner_org_ids'), ('Resources','links_json')]:
        for number, entry in records[name]:
            ref = entry.pop('row_ref')
            if ref not in activities:
                error(name, number, 'row_ref', 'No matching activity row')
                continue
            if name == 'Supporters':
                entry['types'] = [v.strip() for v in entry['types'].split(';') if v.strip()]
            if name == 'Partners':
                activities[ref]['partner_rows'].append(number)
                entry = entry['organization_id']
            if name == 'Resources':
                activities[ref]['resource_rows'].append(number)
                try:
                    L.validate_link(entry)
                except L.LinkValidationError as exc:
                    for key, message in exc.errors.items():
                        error(name, number, key, message)
            else:
                _, invalid = R.validate({field: [entry]})
                invalid_field = 'row_ref'
                if name in ('Contributions', 'Synergies'):
                    invalid_field = ('priority_area' if entry['priority_area'] not in C.PRIORITY_AREAS else
                        'output' if not C.is_valid_output_for_pa(entry['priority_area'], entry['output']) else 'key_activity_code')
                elif name == 'Supporters':
                    invalid_field = 'country' if not R.country_metadata(entry['country']) else 'types'
                elif name == 'Partners':
                    invalid_field = 'organization_id'
                if invalid:
                    invalid_relations.add((ref, field))
                for message in invalid.values():
                    error(name, number, invalid_field, message)
            sources.setdefault(('Activities', activities[ref]['row'], field), location(name, number, 'row_ref'))
            activities[ref]['sources'].setdefault(field, location(name, number, 'row_ref'))
            activities[ref]['data'].setdefault(field, []).append(entry)
            activities[ref].setdefault('relation_sources', {}).setdefault(field, []).append(
                {key: location(name, number, key) for key in SHEETS[name]})
    for ref, item in activities.items():
        data, row = item['data'], item['row']
        for key, message in contradictions(data).items():
            error('Activities', row, key, message)
        invalid_fields = set()
        try:
            F.validate_report_payload(data, True)
        except F.ReportValidationError as exc:
            invalid_fields.update(exc.errors)
            for key, message in exc.errors.items():
                if (ref, key.split('.')[0]) not in invalid_relations:
                    error('Activities', row, key, message)
        try:
            F.validate_report_payload(data, False)
        except F.ReportValidationError as exc:
            for key, message in exc.errors.items():
                if key not in invalid_fields:
                    error('Activities', row, key, message, warnings)
        data['save_as_draft'] = '1'
    if not activities:
        errors.append({'sheet': 'Activities', 'row': 2, 'cell': 'A2', 'field': 'row_ref', 'message': 'Add at least one activity'})
    refs = {(sheet, row): entry.get('row_ref', '') for sheet, entries in original_records.items() for row, entry in entries}
    for issue in errors + warnings:
        issue.setdefault('ref', refs.get((issue.pop('_sheet', ''), issue.pop('_row', 0)), ''))
    return {'version': version, 'rows': [{'ref': ref, **item} for ref, item in activities.items()],
            'errors': errors, 'warnings': warnings}


def contradictions(data):
    """No dejar que un No explícito borre datos proporcionados en el Excel."""
    errors = {}
    extra = {'has_member_state_support': ('supporters_json',), 'has_synergies': ('synergies_json',)}
    for gate, children in F.GATE_RESETS.items():
        if F._gate_answered(data.get(gate)) and not C.normalize_bool(data[gate]):
            for field in tuple(children) + extra.get(gate, ()):
                value = data.get(field)
                if value not in (None, '', 0, '0', [], '[]'):
                    errors[field] = 'This answer contains data, but its Yes/No question is No. Correct the question or clear this answer.'
    return errors


def template_v2(organizations=(), defaults=None, locale='en', example=False):
    from ckanext.theme_ejemplo.ihpix_bulk_workbook import build
    return build(organizations, defaults, locale, example)


def template(organizations=(), defaults=None, locale='en', example=False):
    from ckanext.theme_ejemplo import ihpix_bulk_flat
    from ckanext.theme_ejemplo.ihpix_bulk_workbook import build
    return build(organizations, defaults, locale, example, schema=ihpix_bulk_flat)
