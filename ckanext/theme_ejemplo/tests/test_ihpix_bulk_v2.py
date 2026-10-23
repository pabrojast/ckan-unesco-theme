"""Conservación de campos, errores localizados y compatibilidad de XLSX."""

import io
from functools import lru_cache
import pytest
from openpyxl import load_workbook
from ckanext.theme_ejemplo import (
    ihpix_bulk as B,
    ihpix_bulk_schema as S,
    ihpix_forms as F,
)
from ckanext.theme_ejemplo import ihpix_constants as C
from ckanext.theme_ejemplo.ihpix_bulk_display import present


@lru_cache(maxsize=None)
def template(locale="en", example=False):
    return B.template_v2(locale=locale, example=example)


def edited(changes, locale="en", example=False):
    book = load_workbook(io.BytesIO(template(locale, example)))
    names = {
        r[0].value: r[1].value
        for r in book["_IHP_META"].iter_rows(
            min_row=5, max_row=4 + len(S.SHEET_LABELS), max_col=2
        )
    }
    for section, row, values in changes:
        fields = S.SHEETS[section]
        for field, value in values.items():
            book[names[section]].cell(row, fields.index(field) + 1, value)
    raw = io.BytesIO()
    book.save(raw)
    return raw.getvalue()


@pytest.mark.parametrize("locale", ["en", "es", "fr", "ar"])
def test_example_is_complete_and_preserves_all_relations(locale):
    p = B.parse(template(locale, True))
    assert p["version"] == S.VERSION and p["example"]
    assert not p["errors"] and not p["warnings"]
    assert len(p["rows"]) == 3
    row = p["rows"][2]["data"]
    assert len(row["contributions_json"]) == 2
    assert row["contributions_json"][0]["output"] == "4.1"
    assert row["contributions_json"][1]["output"] == "2.1"
    assert row["synergies_json"][0]["output"] == "1.1"
    assert row["supporters_json"][0]["types"] == ["financial", "hosting"]
    assert len(row["links_json"]) == 2
    assert (
        F.validate_report_payload(p["rows"][1]["data"], True)["stakeholders_knowledge"]
        == 40
    )


def test_every_original_field_is_mapped_exactly_once():
    fields = (
        list(S.SHEETS["Activities"])
        + list(S.SHEETS["Institutions"][1:])
        + list(S.SHEETS["Scope"][1:])
    )
    fields += [v for mapping in S.INDICATORS.values() for v in mapping.values()]
    assert set(fields) == set(B.FIELDS)
    assert len(fields) == len(set(fields)) == 63


def test_common_values_and_external_institution_override_do_not_cross_attribute():
    book = load_workbook(io.BytesIO(template()))
    names = {
        r[0].value: r[1].value
        for r in book["_IHP_META"].iter_rows(
            min_row=5, max_row=4 + len(S.SHEET_LABELS), max_col=2
        )
    }
    start = book[names["Start"]]
    for row, field in enumerate(S.COMMON, 7):
        start.cell(
            row,
            2,
            {
                "institution": "Common",
                "institution_org_id": "common-org",
                "focal_point_name": "Contact",
                "contact_email": "contact@example.org",
                "biennium": C.BIENNIA[0],
            }.get(field, ""),
        )
    activity = book[names["Activities"]]
    activity.cell(5, 2, "One")
    activity.cell(6, 2, "Two")
    institution = book[names["Institutions"]]
    institution.cell(5, 1, "A002")
    institution.cell(5, S.SHEETS["Institutions"].index("institution") + 1, "External")
    raw = io.BytesIO()
    book.save(raw)
    p = B.parse(raw.getvalue())
    assert not p["errors"]
    first, second = [x["data"] for x in p["rows"]]
    assert first["institution_org_id"] == "common-org"
    assert second["institution"] == "External" and not second["institution_org_id"]
    assert second["contact_email"] == "contact@example.org"
    assert p["rows"][1]["sources"]["institution"]["sheet"] == names["Institutions"]


def test_all_indicator_fields_survive_validation_including_zero_and_other():
    changes = [("Activities", 5, {"title": "Full indicator coverage"})]
    expected = {}
    for row, (code, mapping) in enumerate(S.INDICATORS.items(), 5):
        entry = {"row_ref": "A001", "indicator": code, "answer": "yes"}
        for column, field in mapping.items():
            if column == "answer":
                expected[field] = "yes"
                continue
            val = {
                "quantity": "10",
                "youth": "0",
                "female": "7",
                "other": "Other detail",
                "group_names": "Group One",
            }.get(column)
            if column == "types":
                val = ";".join(F.MULTI_FIELDS[field][:2])
            entry[column] = val
            expected[field] = val.split(";") if field in F.MULTI_FORM_FIELDS else val
        changes.append(("Indicators", row, entry))
    p = B.parse(edited(changes))
    assert not p["errors"]
    actual = F.validate_report_payload(p["rows"][0]["data"], True)
    expected = F.validate_report_payload(
        dict(expected, title="Full indicator coverage"), True
    )
    for mapping in S.INDICATORS.values():
        for field in mapping.values():
            assert actual[field] == expected[field], field
    display = present({"preview": p}, "en")
    assert any(
        v["value"] == "0"
        for section in display["rows"][0]["sections"]
        for v in section["fields"]
    )


def test_negative_integer_points_to_indicator_cell_and_is_not_a_warning():
    p = B.parse(
        edited(
            [
                ("Activities", 5, {"title": "Training"}),
                (
                    "Indicators",
                    8,
                    {
                        "row_ref": "A001",
                        "indicator": "2",
                        "answer": "yes",
                        "quantity": "-3",
                    },
                ),
            ]
        )
    )
    issue = next(e for e in p["errors"] if e["field"] == "stakeholders_knowledge")
    assert (issue["sheet"], issue["cell"], issue["ref"]) == ("Indicators", "D8", "A001")
    assert not any(w["field"] == "stakeholders_knowledge" for w in p["warnings"])


@pytest.mark.parametrize("answer", ["no", "false", "0"])
def test_no_cannot_silently_discard_a_count(answer):
    p = B.parse(
        edited(
            [
                ("Activities", 5, {"title": "Contradiction"}),
                (
                    "Indicators",
                    5,
                    {
                        "row_ref": "A001",
                        "indicator": "2",
                        "answer": answer,
                        "quantity": "25",
                    },
                ),
            ]
        )
    )
    assert any(
        e["field"] == "stakeholders_knowledge" and "No" in e["message"]
        for e in p["errors"]
    )


def test_no_cannot_discard_supporters_or_synergies():
    p = B.parse(
        edited(
            [("Scope", 7, {"has_member_state_support": "no", "has_synergies": "no"})],
            example=True,
        )
    )
    assert {"supporters_json", "synergies_json"} <= {e["field"] for e in p["errors"]}


def test_duplicate_indicator_and_irrelevant_columns_are_rejected():
    p = B.parse(
        edited(
            [
                ("Activities", 5, {"title": "One"}),
                (
                    "Indicators",
                    5,
                    {
                        "row_ref": "A001",
                        "indicator": "4",
                        "answer": "yes",
                        "quantity": "2",
                        "youth": "1",
                    },
                ),
                (
                    "Indicators",
                    6,
                    {
                        "row_ref": "A001",
                        "indicator": "4",
                        "answer": "yes",
                        "quantity": "3",
                    },
                ),
            ]
        )
    )
    assert any("does not apply" in e["message"] for e in p["errors"])
    assert any("one row" in e["message"] for e in p["errors"])


def test_orphans_duplicates_formulas_and_blank_prepared_references():
    assert len(B.parse(template())["rows"]) == 0
    p = B.parse(
        edited(
            [
                ("Activities", 5, {"title": "One"}),
                ("Activities", 6, {"row_ref": "A001", "title": "Duplicate"}),
                ("Activities", 7, {"title": "=1+1"}),
                (
                    "Supporters",
                    5,
                    {"row_ref": "A999", "country": "Chile", "financial": "yes"},
                ),
            ]
        )
    )
    assert len(p["rows"]) == 2
    assert any("unique" in e["message"] for e in p["errors"])
    assert any("matching" in e["message"] for e in p["errors"])
    assert any("Formulas" in e["message"] for e in p["errors"])


def test_complete_text_dates_scope_and_resources_survive():
    changes = [
        (
            "Activities",
            5,
            {
                "title": "All text",
                "start_date": "2026-01-01",
                "end_date": "2026-01-02",
                "description": "Description",
                "outcomes": "Results",
                "priority_area": "PA1",
                "output": "1.10",
                "key_activity_code": "1.10.1",
                "key_activity": "Contribution text",
                "link": "https://example.org",
            },
        ),
        (
            "Institutions",
            5,
            dict(
                row_ref="A001",
                **{
                    f: {
                        "focal_point_name": "Person",
                        "contact_email": "person@example.org",
                        "institution": "Institute",
                        "institution_org_id": "verified-org",
                        "institution_type": C.LEAD_INSTITUTION_TYPES[0],
                        "institution_type_other": "Other category",
                        "country": "CL",
                        "biennium": C.BIENNIA[0],
                        "partners": "External partner",
                        "unesco_secretariat_participation": "yes",
                    }[f]
                    for f in S.SHEETS["Institutions"][1:]
                }
            ),
        ),
        (
            "Scope",
            5,
            {
                "row_ref": "A001",
                "has_member_state_support": "yes",
                "supporting_member_state": "Kenya",
                "has_flagship": "yes",
                "flagships": C.FLAGSHIPS[0],
                "cross_cutting_wg": C.CROSS_CUTTING_WGS[0],
                "has_synergies": "yes",
                "synergies": "Synergy description",
                "regions_benefit": "yes",
                "regions": C.REGIONS[0],
                "member_states": "Chile;Kenya",
                "additional_notes": "All notes",
                "image_url": "https://example.org/image.png",
            },
        ),
        ("Partners", 5, {"row_ref": "A001", "organization_id": "partner-org"}),
        (
            "Resources",
            5,
            {
                "row_ref": "A001",
                "link_type": "event",
                "title": "Event",
                "url": "https://example.org/event",
                "description": "Resource detail",
                "event_date": "2026-01-01",
                "target_kind": "url",
            },
        ),
    ]
    p = B.parse(edited(changes))
    assert not p["errors"]
    row = p["rows"][0]["data"]
    saved = F.validate_report_payload(row, True)
    assert row["partner_org_ids"] == ["partner-org"]
    assert row["links_json"][0]["description"] == "Resource detail"
    for field in (
        S.SHEETS["Activities"][1:]
        + S.SHEETS["Institutions"][1:]
        + S.SHEETS["Scope"][1:]
    ):
        assert row[field] not in ("", None), field
        assert saved[field] not in ("", None), field


def test_resource_error_points_to_actual_v2_column():
    p = B.parse(
        edited(
            [
                ("Activities", 5, {"title": "One"}),
                (
                    "Resources",
                    9,
                    {
                        "row_ref": "A001",
                        "link_type": "publication",
                        "title": "Paper",
                        "url": "not-a-url",
                    },
                ),
            ]
        )
    )
    issue = next(e for e in p["errors"] if e["field"] == "url")
    assert issue["cell"] == "D9" and issue["ref"] == "A001"


def test_preview_pagination_is_bounded():
    p = B.parse(
        edited([("Activities", 5 + i, {"title": "Activity %s" % i}) for i in range(42)])
    )
    display = present({"preview": p}, page=2)
    assert display["pages"] == 3 and len(display["rows"]) == 20
    assert display["rows"][0]["ref"] == "A021"


def test_stale_dimensions_do_not_hide_activities_or_relations():
    import zipfile, re

    raw = edited(
        [
            ("Activities", 5, {"title": "One"}),
            ("Activities", 6, {"title": "Two"}),
            (
                "Resources",
                8,
                {
                    "row_ref": "A002",
                    "link_type": "publication",
                    "title": "Paper",
                    "url": "https://example.org",
                },
            ),
        ]
    )
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as source, zipfile.ZipFile(
        output, "w"
    ) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename.startswith("xl/worksheets/"):
                data = re.sub(
                    rb'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"', data
                )
            target.writestr(info, data)
    parsed = B.parse(output.getvalue())
    assert not parsed["errors"]
    assert [r["data"]["title"] for r in parsed["rows"]] == ["One", "Two"]
    assert parsed["rows"][1]["data"]["links_json"][0]["title"] == "Paper"


def test_reordered_common_labels_cannot_misattribute_values():
    book = load_workbook(io.BytesIO(template()))
    book["Activities"]["B5"] = "One"
    sheet = book["Start here"]
    sheet["A7"], sheet["A9"] = sheet["A9"].value, sheet["A7"].value
    raw = io.BytesIO()
    book.save(raw)
    with pytest.raises(ValueError, match="Common reporting rows"):
        B.parse(raw.getvalue())


def test_missing_optional_rows_do_not_point_to_another_activity():
    parsed = B.parse(
        edited(
            [
                ("Activities", 5, {"title": "One"}),
                ("Activities", 6, {"title": "Two"}),
                (
                    "Indicators",
                    5,
                    {
                        "row_ref": "A001",
                        "indicator": "2",
                        "answer": "yes",
                        "quantity": "3",
                    },
                ),
            ]
        )
    )
    warning = next(
        w
        for w in parsed["warnings"]
        if w["ref"] == "A002" and w["field"] == "kpi_2_active"
    )
    assert (
        warning["missing_row"] and warning["cell"] == "" and warning["indicator"] == "2"
    )


def test_legacy_persisted_warnings_are_still_visible():
    book = load_workbook(io.BytesIO(B.template_v1()))
    book["Activities"].append(["Old", "Old draft"])
    raw = io.BytesIO()
    book.save(raw)
    p = B.parse(raw.getvalue())
    for w in p["warnings"]:
        w.pop("ref", None)
    display = present({"preview": p})
    assert display["incomplete"] == 1
    assert len(display["rows"][0]["warnings"]) == len(p["warnings"])


def test_relation_errors_are_not_repeated_as_activity_errors():
    p = B.parse(
        edited(
            [
                (
                    "Activities",
                    5,
                    {"title": "One", "priority_area": "PA1", "output": "1.1"},
                ),
                (
                    "Contributions",
                    9,
                    {
                        "row_ref": "A001",
                        "relation": "contribution",
                        "priority_area": "PA1",
                        "output": "2.1",
                    },
                ),
            ]
        )
    )
    assert len(p["errors"]) == 1
    assert p["errors"][0]["cell"] == "D9"


def test_additional_contribution_never_silently_becomes_primary():
    p = B.parse(
        edited(
            [
                ("Activities", 5, {"title": "One"}),
                (
                    "Contributions",
                    9,
                    {
                        "row_ref": "A001",
                        "relation": "contribution",
                        "priority_area": "PA2",
                        "output": "2.1",
                    },
                ),
            ]
        )
    )
    assert any(e["field"] == "output" and e["cell"] == "H5" for e in p["errors"])


def test_dates_have_native_workbook_validation():
    book = load_workbook(io.BytesIO(template()))
    for name, cells in [
        ("Activities", ["C5", "D5"]),
        ("Publications and resources", ["F5"]),
    ]:
        for cell in cells:
            assert any(
                cell in rule.sqref
                and rule.type == "custom"
                and "DATE(" in rule.formula1
                for rule in book[name].data_validations.dataValidation
            )


@pytest.mark.parametrize("locale", ["es", "fr", "ar"])
def test_localized_example_headings_and_indicator_type_choices(locale):
    book = load_workbook(io.BytesIO(template(locale, True)))
    tr = S.translator(locale)
    assert book[tr(S.SHEET_LABELS["Activities"])]["B4"].value == tr(S.LABELS["title"])
    indicator = book[tr(S.SHEET_LABELS["Indicators"])]
    assert indicator["G5"].value == tr(C.KNOWLEDGE_PRODUCT_TYPES[0])
    reference = book[tr(S.SHEET_LABELS["Reference"])]
    assert reference["A2"].value == tr("Yes or no")


@pytest.mark.parametrize("malformation", ["reversed_rows", "duplicate_row", "duplicate_cell", "wrong_row"])
def test_malformed_xml_never_silently_drops_or_overwrites_data(malformation):
    import re, zipfile
    raw = edited([("Activities", 5, {"title": "One"}), ("Activities", 6, {"title": "Two"})])
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as source, zipfile.ZipFile(output, "w") as target:
        for info in source.infolist():
            value = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet4.xml":
                text = value.decode()
                first = re.search(r'<row r="5"[^>]*>.*?</row>', text).group()
                second = re.search(r'<row r="6"[^>]*>.*?</row>', text).group()
                if malformation == "reversed_rows":
                    text = text.replace(first + second, second + first)
                elif malformation == "duplicate_row":
                    text = text.replace(first, first + first)
                elif malformation == "duplicate_cell":
                    cell = re.search(r'<c r="B5"[^>]*>.*?</c>', first).group()
                    text = text.replace(cell, cell + cell)
                else:
                    text = text.replace('r="B5"', 'r="B6"')
                value = text.encode()
            target.writestr(info, value)
    with pytest.raises(ValueError, match="Worksheet structure changed"):
        B.parse(output.getvalue())
