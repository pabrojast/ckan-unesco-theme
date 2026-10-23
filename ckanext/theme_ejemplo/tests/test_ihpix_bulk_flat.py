"""Pruebas del uso plano: copiar filas, ordenar, títulos y conservación de campos."""

import io
from functools import lru_cache
import pytest
from openpyxl import load_workbook
from ckanext.theme_ejemplo import ihpix_bulk as B, ihpix_bulk_flat as S
from ckanext.theme_ejemplo.ihpix_bulk_display import present


@lru_cache(maxsize=None)
def workbook(locale="en", example=False):
    return B.template(locale=locale, example=example)


def save(book):
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def names(book):
    return {
        r[0].value: r[1].value
        for r in book["_IHP_META"].iter_rows(
            min_row=5, max_row=4 + len(S.SHEET_LABELS), max_col=2
        )
    }


@pytest.mark.parametrize("locale", ["en", "es", "fr", "ar"])
def test_flat_example_preserves_every_value_from_v2(locale):
    flat = B.parse(workbook(locale, True))
    old = B.parse(B.template_v2(locale=locale, example=True))
    assert flat["version"] == S.VERSION
    assert not flat["errors"] and not flat["warnings"]
    assert [r["data"] for r in flat["rows"]] == [r["data"] for r in old["rows"]]
    book = load_workbook(io.BytesIO(workbook(locale, True)))
    assert book.active.title == names(book)["Activities"]
    assert book.sheetnames[0] == book.active.title
    assert book.active["A5"].value.startswith(S.translator(locale)("[EXAMPLE]"))
    assert set(S.SHEETS["Activities"]) == set(B.FIELDS) - {"row_ref"}
    assert len(S.SHEETS["Activities"]) == len(set(S.SHEETS["Activities"])) == 62
    assert not any(
        key in names(book) for key in ("Institutions", "Scope", "Indicators")
    )
    assert all(
        not value.get("missing_row") for value in flat["rows"][0]["sources"].values()
    )


def test_blank_template_has_no_personal_or_example_data_to_remove():
    book = load_workbook(
        io.BytesIO(
            B.template(
                defaults={
                    "institution": "My institution",
                    "contact_email": "me@example.org",
                }
            )
        )
    )
    assert all(c.value is None for row in book.active.iter_rows(min_row=5) for c in row)
    result = B.parse(save(book))
    assert len(result["errors"]) == 1
    assert result["errors"][0]["cell"] == "A5"


def test_title_only_creates_a_valid_draft_without_reference_codes():
    book = load_workbook(io.BytesIO(workbook()))
    book.active["A5"] = "My activity"
    result = B.parse(save(book))
    assert not result["errors"]
    assert result["rows"][0]["data"]["title"] == "My activity"
    assert result["warnings"]
    assert all(x["sheet"] == "Activities" and x["cell"] for x in result["warnings"])
    assert present({"preview": result})["flat"]


def test_copying_and_sorting_rows_preserves_optional_details_by_title():
    book = load_workbook(io.BytesIO(workbook(example=True)))
    main = book.active
    first = [c.value for c in main[5]]
    last = [c.value for c in main[7]]
    for col, value in enumerate(last, 1):
        main.cell(5, col).value = value
    for col, value in enumerate(first, 1):
        main.cell(7, col).value = value
    for col, value in enumerate(first, 1):
        main.cell(8, col).value = value
    main.cell(8, 1, "Copied activity")
    result = B.parse(save(book))
    assert not result["errors"]
    assert len(result["rows"]) == 4
    assert result["rows"][0]["data"]["title"] == last[0]
    assert len(result["rows"][0]["data"]["links_json"]) == 2
    assert not result["rows"][2]["data"].get("links_json")
    assert result["rows"][3]["data"]["title"] == "Copied activity"


@pytest.mark.parametrize("case", ["missing", "duplicate"])
def test_extra_details_never_attach_to_an_ambiguous_or_unknown_title(case):
    book = load_workbook(io.BytesIO(workbook(example=True)))
    if case == "missing":
        book.active["A7"] = "Renamed activity"
    else:
        book.active["A5"] = book.active["A7"].value
    result = B.parse(save(book))
    assert any(
        e["field"] == "row_ref" and e["cell"] == "A5" and "exactly once" in e["message"]
        for e in result["errors"]
    )
    display = present({"preview": result})
    assert all(
        e["label"] == "Activity title"
        for e in display["errors"]
        if e["field"] == "row_ref"
    )


def test_zero_no_and_invalid_counts_use_the_actual_flat_cells():
    book = load_workbook(io.BytesIO(workbook(example=True)))
    fields = S.SHEETS["Activities"]
    quantity = fields.index("num_knowledge_products") + 1
    book.active.cell(5, quantity, -1)
    result = B.parse(save(book))
    issue = next(e for e in result["errors"] if e["field"] == "num_knowledge_products")
    assert issue["cell"] == book.active.cell(5, quantity).coordinate
    book.active.cell(5, quantity, 0)
    result = B.parse(save(book))
    assert not result["errors"]
    assert result["rows"][0]["data"]["num_knowledge_products"] == "0"
    assert result["rows"][0]["data"]["kpi_4_active"] == "no"


def test_native_validations_and_title_dropdowns_need_no_formulas_in_input_cells():
    book = load_workbook(io.BytesIO(workbook(example=True)))
    fields = S.SHEETS["Activities"]
    for field, kind in [
        ("num_knowledge_products", "whole"),
        ("start_date", "custom"),
        ("kpi_1a_active", "list"),
        ("member_states", "list"),
    ]:
        cell = book.active.cell(5, fields.index(field) + 1).coordinate
        assert any(
            cell in rule.sqref and rule.type == kind
            for rule in book.active.data_validations.dataValidation
        )
    for key in ("Contributions", "Supporters", "Partners", "Resources"):
        sheet = book[names(book)[key]]
        assert sheet["A4"].value == "Activity title"
        assert any(
            "A5" in rule.sqref and rule.formula1 == "ihp_activity_refs"
            for rule in sheet.data_validations.dataValidation
        )
    assert all(c.data_type != "f" for sheet in book for row in sheet for c in row)


@pytest.mark.parametrize("locale,heading", [("es","Actividad"),("fr","Activité"),("ar","النشاط")])
def test_main_group_heading_is_localized(locale,heading):
    book=load_workbook(io.BytesIO(workbook(locale,True)))
    assert book.active["A2"].value==heading
