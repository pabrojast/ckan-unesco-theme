# -*- coding: utf-8 -*-
"""Libro guiado v2 y adaptación al validador de importación existente."""
import datetime
import io
import textwrap
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from ckanext.theme_ejemplo import ihpix_bulk_schema as S
from ckanext.theme_ejemplo import (
    ihpix_forms as F,
    ihpix_constants as C,
    ihpix_relations as R,
    ihpix_links as L,
)


def _(text):
    return text


GUIDE = _(
    "Start with Activities. Complete the other sheets only when relevant. One activity per row; keep its reference when adding details on another sheet."
)
OPTIONAL_GUIDE = _(
    "Optional. Select the activity reference from Activities. Leave unanswered questions blank; No is an explicit answer."
)
MULTI_HELP = _(
    "For multiple choices, separate names with ;. See Reference lists. Example: Chile; Kenya."
)
DATE_HELP = _("Use YYYY-MM-DD, for example 2026-09-15.")
NUMBER_HELP = _("Enter a whole number from 0 to 2147483647. Leave blank if unanswered.")
CODE_HELP = _(
    "Choose a value from the list. The Output must belong to its Priority Area, and the key activity to its Output."
)
COMMON_HELP = _(
    "Used for every activity unless an exception is entered in Institutions."
)


def choices(organizations, locale, tr):
    result = OrderedDict()
    result["yesno"] = [("yes", tr("Yes")), ("no", tr("No"))]
    result["priority_area"] = [
        (k, k + " · " + tr(v)) for k, v in C.PRIORITY_AREAS.items()
    ]
    short = lambda text: text if len(text) <= 85 else text[:82].rstrip() + "…"
    result["output"] = [
        (k, k + " · " + short(tr(C.output_title(k)))) for k in C.all_output_codes()
    ]
    result["key_activity_code"] = [
        (x["code"], x["code"] + " · " + short(tr(x["title"])))
        for entries in C.load_key_activities().values()
        for x in entries
    ]
    result["country"] = [
        (
            x["country_title_en"],
            x.get("country_title_" + locale) or x["country_title_en"],
        )
        for x in R.countries()
    ]
    result["organization_id"] = [
        (o["id"], (o.get("title") or o["name"]) + " · " + o.get("name", o["id"]))
        for o in organizations
    ]
    result["biennium"] = [(v, v) for v in C.BIENNIA]
    result["institution_type"] = [(v, tr(v)) for v in C.LEAD_INSTITUTION_TYPES]
    for field, values in F.MULTI_FIELDS.items():
        if values:
            result[field] = [
                (v, tr(C.flagship_label(v) if field == "flagships" else v))
                for v in values
            ]
    result["indicator"] = [
        (key, key + " · " + tr(value)) for key, value in S.INDICATOR_LABELS.items()
    ]
    result["relation"] = [
        ("contribution", tr("Additional contribution")),
        ("synergy", tr("Synergy")),
    ]
    result["link_type"] = [(key, tr(value)) for key, value in L.TYPE_LABELS.items()]
    result["target_kind"] = [
        ("url", tr("External link")),
        ("package", tr("Dataset or document")),
        ("page", tr("Event page")),
        ("course", tr("Course")),
    ]
    return result


def choice_field(field):
    if field in F.BOOL_FIELDS or field in ("answer", "financial", "in_kind", "hosting"):
        return "yesno"
    return {
        "institution_org_id": "organization_id",
        "supporting_member_state": "country",
        "member_states": "country",
    }.get(field, field)


def field_help(field, tr, section=None):
    if field in F.MULTI_FORM_FIELDS or field == "types":
        return tr(MULTI_HELP)
    if field in ("start_date", "end_date", "event_date"):
        return tr(DATE_HELP)
    if field in ("quantity", "youth", "female") or field in F.INT_FIELDS:
        return tr(NUMBER_HELP)
    if field in ("priority_area", "output", "key_activity_code"):
        return tr(CODE_HELP)
    if section == "Resources" and field == "description":
        return tr("Maximum characters:") + " " + str(L.DESCRIPTION_MAX)
    maximum = (
        F.SHORT_TEXT_MAX
        if field in ("description", "outcomes")
        else F.LONG_TEXT_MAX.get(field)
    )
    if maximum:
        return tr("Maximum characters:") + " " + str(maximum)
    if field == "row_ref":
        return tr("Use the same activity reference on every sheet, for example A001.")
    return tr("Leave blank if unanswered.")


def example_records(tr):
    base = {
        "institution": tr("Example Water Institute"),
        "focal_point_name": tr("Example coordinator"),
        "contact_email": "example@example.org",
        "institution_type": "University",
        "country": "Chile",
        "biennium": C.BIENNIA[0],
    }
    records = {name: [] for name in S.SHEETS}
    for i, (name, output) in enumerate(
        [
            ("River monitoring study", "1.1"),
            ("Water management training", "2.1"),
            ("Regional water cooperation", "4.1"),
        ],
        1,
    ):
        key = C.load_key_activities()[output][0]["code"]
        records["Activities"].append(
            {
                "row_ref": "A%03d" % i,
                "title": tr("[EXAMPLE]") + " " + tr(name),
                "start_date": "2026-09-01",
                "end_date": "2026-09-15",
                "description": tr(
                    "Fictional activity for learning how to complete this workbook."
                ),
                "outcomes": tr(
                    "Example results only. Replace these with the results of your activity."
                ),
                "priority_area": C.priority_area_for_output(output),
                "output": output,
                "key_activity_code": key,
                "key_activity": tr(
                    "Explain how the activity contributed to the selected key activity."
                ),
                "link": "https://example.org/activity",
            }
        )
        ref = "A%03d" % i
        records["Institutions"].append(
            {"row_ref": ref, "unesco_secretariat_participation": "no"}
        )
        records["Scope"].append(
            {
                "row_ref": ref,
                "has_member_state_support": "yes" if i == 3 else "no",
                "has_flagship": "no",
                "has_synergies": "yes" if i == 3 else "no",
                "regions_benefit": "yes",
                "regions": (
                    "Latin America and the Caribbean;Africa"
                    if i == 3
                    else "Latin America and the Caribbean"
                ),
                "member_states": "Chile;Kenya" if i == 3 else "Chile",
                "additional_notes": tr(
                    "Fictional example. Do not submit as a real activity."
                ),
            }
        )
        for key in S.INDICATORS:
            row = {"row_ref": ref, "indicator": key, "answer": "no"}
            if i == 1 and key == "1a":
                row.update(answer="yes", quantity=1, types=C.KNOWLEDGE_PRODUCT_TYPES[0])
            if i == 2 and key == "2":
                row.update(
                    answer="yes",
                    quantity=40,
                    youth=15,
                    female=22,
                    types=C.KNOWLEDGE_ACTIVITY_TYPES[0],
                )
            if i == 2 and key == "3":
                row.update(answer="yes", quantity=2, types=C.TRAINING_TYPES[0])
            if i == 3 and key == "6":
                row.update(answer="yes", quantity=2)
            records["Indicators"].append(row)
    records["Contributions"] = [
        {
            "row_ref": "A003",
            "relation": "contribution",
            "priority_area": "PA2",
            "output": "2.1",
            "key_activity_code": C.load_key_activities()["2.1"][0]["code"],
        },
        {
            "row_ref": "A003",
            "relation": "synergy",
            "priority_area": "PA1",
            "output": "1.1",
            "key_activity_code": "",
        },
    ]
    records["Supporters"] = [
        {
            "row_ref": "A003",
            "country": "Kenya",
            "financial": "yes",
            "in_kind": "no",
            "hosting": "yes",
        }
    ]
    records["Resources"] = [
        {
            "row_ref": "A003",
            "link_type": "publication",
            "title": tr("Example publication"),
            "url": "https://example.org/publication",
            "description": tr(
                "Fictional resource used to illustrate an external link."
            ),
            "target_kind": "url",
        },
        {
            "row_ref": "A003",
            "link_type": "event",
            "title": tr("Example workshop"),
            "url": "https://example.org/workshop",
            "event_date": "2026-09-10",
        },
    ]
    return base, records


def build(organizations=(), defaults=None, locale="en", example=False, schema=S):
    S = schema
    flat = getattr(S, "FLAT", False)
    from ckanext.theme_ejemplo.ihpix_bulk import MAX_ROWS

    locale = locale if locale in ("en", "es", "fr", "ar") else "en"
    tr = S.translator(locale)
    book = Workbook()
    book.remove(book.active)
    names = {key: tr(value)[:31] for key, value in S.SHEET_LABELS.items()}
    meta = book.create_sheet("_IHP_META")
    meta.append(["version", S.VERSION])
    meta.append(["locale", locale])
    meta.append(["example", bool(example)])
    meta.append(["sheet", "name"])
    for key, name in names.items():
        meta.append([key, name])
    for row, field in enumerate(S.COMMON, 20):
        meta.cell(row, 1, field)
        meta.cell(row, 2, S.label(field, tr))
    meta.sheet_state = "hidden"
    catalog = choices(organizations, locale, tr)
    lookup = {
        field: {str(value): display for value, display in items}
        for field, items in catalog.items()
    }
    opts = book.create_sheet("_IHP_OPTIONS")
    opts.append(["field", "label", "value"])
    opts.sheet_state = "hidden"
    ranges = {}
    for field, items in catalog.items():
        start = opts.max_row + 1
        for value, display in items:
            opts.append([field, display, value])
        if items:
            name = "ihp_" + field
            definition = DefinedName(
                name, attr_text="'_IHP_OPTIONS'!$B$%s:$B$%s" % (start, opts.max_row)
            )
            if hasattr(book.defined_names, "add"):
                book.defined_names.add(definition)
            else:
                book.defined_names.append(definition)
            ranges[field] = name

    def display(field, value):
        category = choice_field(field)
        if field in F.MULTI_FORM_FIELDS or field == "types":
            return "; ".join(
                lookup.get(category, {}).get(v.strip(), v.strip())
                for v in str(value).split(";")
            )
        return lookup.get(category, {}).get(str(value), value)

    def validate(sheet, field, address):
        category = choice_field(field)
        rule = None
        if field in ("start_date", "end_date", "event_date"):
            # La regla acepta fechas de Excel y texto ISO, sin fórmulas en las celdas.
            cell = address.split(":")[0]
            formula = 'OR(AND(ISNUMBER({0}),{0}>=DATE(1900,1,1),{0}<=DATE(9999,12,31)),IFERROR(TEXT(DATE(LEFT({0},4),MID({0},6,2),RIGHT({0},2)),"yyyy-mm-dd")={0},FALSE))'.format(
                cell
            )
            rule = DataValidation(type="custom", formula1=formula, allow_blank=True)
        elif field in ("quantity", "youth", "female") or field in F.INT_FIELDS:
            rule = DataValidation(
                type="whole",
                operator="between",
                formula1=0,
                formula2=2147483647,
                allow_blank=True,
            )
        elif (
            category in ranges
            and (flat or field not in F.MULTI_FORM_FIELDS)
            and field != "types"
        ):
            rule = DataValidation(
                type="list", formula1=ranges[category], allow_blank=True
            )
        elif field == "row_ref":
            rule = DataValidation(
                type="list", formula1="ihp_activity_refs", allow_blank=True
            )
        if rule:
            rule.showErrorMessage = not (flat and field in F.MULTI_FORM_FIELDS)
            rule.errorTitle = tr("Check this value")
            rule.error = tr(
                DATE_HELP
                if field in ("start_date", "end_date", "event_date")
                else "Choose an available option or enter a valid number."
            )
            rule.showInputMessage = True
            rule.promptTitle = (
                tr("Activity title")
                if flat and field == "row_ref"
                else S.label(field, tr)
            )[:32]
            rule.prompt = (
                tr("Choose the activity title from Activities.")
                if flat and field == "row_ref"
                else field_help(field, tr)
            )[:255]
            sheet.add_data_validation(rule)
            rule.add(address)

    blue = "005A96"
    light = "EAF3FA"
    gray = "536579"
    start = book.create_sheet(names["Start"])
    start.append([tr("IHP-IX · Import activities")])
    start.merge_cells("A1:D1")
    start.append([tr(GUIDE)])
    start.merge_cells("A2:D3")
    start["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    start.row_dimensions[2].height = 42
    start["A5"] = tr("1. Fill in the common reporting details")
    start.merge_cells("A5:D5")
    start["A6"] = tr("Field")
    start["B6"] = tr("Shared value")
    start["C6"] = tr("Help")
    base = dict(defaults or {})
    example_data = {}
    if example:
        base, example_data = example_records(tr)
        if flat:
            combined = {
                entry["row_ref"]: dict(base, **entry)
                for entry in example_data["Activities"]
            }
            for key in ("Institutions", "Scope"):
                for entry in example_data[key]:
                    combined[entry["row_ref"]].update(entry)
            for entry in example_data["Indicators"]:
                for column, field in S.INDICATORS[entry["indicator"]].items():
                    combined[entry["row_ref"]][field] = entry.get(column, "")
            for key in ("Contributions", "Supporters", "Partners", "Resources"):
                for entry in example_data[key]:
                    entry["row_ref"] = combined[entry["row_ref"]]["title"]
            example_data["Activities"] = list(combined.values())
    for row, field in enumerate(S.COMMON, 7):
        start.cell(row, 1, S.label(field, tr))
        start.cell(row, 2, display(field, base.get(field, "")))
        start.cell(row, 3, field_help(field, tr))
        start.merge_cells(start_row=row, start_column=3, end_row=row, end_column=4)
        start.row_dimensions[row].height = 36
        for col in (1, 2, 3):
            start.cell(row, col).alignment = Alignment(
                wrap_text=True, vertical="center"
            )
        start.cell(row, 2).fill = PatternFill("solid", fgColor=light)
        start.cell(row, 2).number_format = "@"
        validate(start, field, "B" + str(row))
    start["A15"] = tr(COMMON_HELP)
    start.merge_cells("A15:D15")
    start["A15"].alignment = Alignment(wrap_text=True)
    start.row_dimensions[15].height = 30
    start["A17"] = tr("2. Enter activities, then add details only where needed")
    start.merge_cells("A17:D17")
    for row, key in enumerate(S.SHEETS, 19):
        cell = start.cell(row, 1, names[key])
        cell.hyperlink = "#" + quote_sheetname(names[key]) + "!A1"
        cell.style = "Hyperlink"
        start.cell(
            row,
            2,
            (
                tr("Start here: one activity per row.")
                if key == "Activities"
                else tr("Optional: complete only when relevant.")
            ),
        )
    row = 29
    for help_text in [
        MULTI_HELP,
        _(
            "Only the activity title is required to create a draft. Other required answers can be completed before submission."
        ),
        _(
            "3. Upload this workbook, review the results and confirm creation of drafts."
        ),
        _("Files and publications can also be added after creating the drafts."),
        _(
            "The example contains fictional activities. Use the empty template for real reports."
        ),
    ]:
        start.cell(row, 1, tr(help_text))
        start.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        start.cell(row, 1).alignment = Alignment(wrap_text=True)
        start.row_dimensions[row].height = 30
        row += 1
    for col, width in [("A", 39), ("B", 48), ("C", 43), ("D", 5)]:
        start.column_dimensions[col].width = width
    start.freeze_panes = "B7"
    for row in (5, 17):
        start.cell(row, 1).font = Font(name="Calibri", bold=True, size=13, color=blue)
        start.row_dimensions[row].height = 32
    for cell in start[6][:3]:
        cell.font = Font(name="Calibri", bold=True, color=blue)
    for row in range(19, 27):
        start.row_dimensions[row].height = 24
    refs = DefinedName(
        "ihp_activity_refs",
        attr_text=quote_sheetname(names["Activities"])
        + "!$A$5:$A$"
        + str(MAX_ROWS + 4),
    )
    if hasattr(book.defined_names, "add"):
        book.defined_names.add(refs)
    else:
        book.defined_names.append(refs)
    for key, fields in S.SHEETS.items():
        sheet = book.create_sheet(names[key])
        sheet.append(list(fields))
        sheet.row_dimensions[1].hidden = True
        sheet.cell(2, 1, names[key])
        sheet.merge_cells(
            start_row=2, start_column=1, end_row=2, end_column=len(fields)
        )
        sheet.cell(
            3,
            1,
            tr(
                (S.GUIDE if key == "Activities" else S.EXTRA_GUIDE)
                if flat
                else (GUIDE if key == "Activities" else OPTIONAL_GUIDE)
            ),
        )
        sheet.merge_cells(
            start_row=3, start_column=1, end_row=3, end_column=len(fields)
        )
        sheet.cell(3, 1).alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[3].height = 42
        sheet.freeze_panes = "B5"
        limit = MAX_ROWS * 8 if key == "Indicators" else MAX_ROWS
        for col, field in enumerate(fields, 1):
            visible_field = (
                "resource_" + field
                if key == "Resources" and field in ("title", "description")
                else field
            )
            heading = S.label(visible_field, tr)
            if flat and field == "row_ref":
                heading = tr("Activity title")
            if flat and key == "Activities":
                for mapping in S.INDICATORS.values():
                    for short, actual in mapping.items():
                        if actual == field:
                            heading = tr(S.LABELS[short])
            cell = sheet.cell(4, col, heading)
            cell.comment = Comment(
                (
                    tr("Choose the activity title from Activities.")
                    if flat and field == "row_ref"
                    else field_help(field, tr, key)
                ),
                "IHP-IX",
            )
            cell.fill = PatternFill("solid", fgColor=blue)
            cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            letter = get_column_letter(col)
            sheet.column_dimensions[letter].number_format = (
                "0"
                if field in ("quantity", "youth", "female") or field in F.INT_FIELDS
                else (
                    "yyyy-mm-dd"
                    if field in ("start_date", "end_date", "event_date")
                    else "@"
                )
            )
            sheet.column_dimensions[letter].width = 18 if field == "row_ref" else 32
            if field in (
                "description",
                "outcomes",
                "key_activity",
                "additional_notes",
                "url",
            ):
                sheet.column_dimensions[letter].width = 46
            if key != "Activities" or field != "row_ref":
                validate(sheet, field, letter + "5:" + letter + str(limit + 4))
            for row in range(5, 35):
                c = sheet.cell(row, col)
                c.number_format = (
                    "0"
                    if field in ("quantity", "youth", "female") or field in F.INT_FIELDS
                    else (
                        "yyyy-mm-dd"
                        if field in ("start_date", "end_date", "event_date")
                        else "@"
                    )
                )
                c.alignment = Alignment(wrap_text=True, vertical="top")
                c.fill = PatternFill("solid", fgColor="F4F8FC" if row % 2 else "FFFFFF")
        sheet.row_dimensions[4].height = 42
        for row in range(5, 35):
            sheet.row_dimensions[row].height = (
                54 if key in ("Activities", "Scope", "Resources") else 36
            )
        sheet.auto_filter.ref = "A4:" + get_column_letter(len(fields)) + str(limit + 4)
        if key == "Activities" and not flat:
            for i in range(MAX_ROWS):
                sheet.cell(i + 5, 1, "A%03d" % (i + 1))
        if key == "Resources":
            sheet.column_dimensions.group("G", "H", hidden=True)
        sheet.sheet_properties.tabColor = blue if key == "Activities" else "90A9BF"
        for row, entry in enumerate(example_data.get(key, []), 5):
            lines = 1
            for col, field in enumerate(fields, 1):
                actual_field = (
                    S.INDICATORS.get(entry.get("indicator"), {}).get(field, field)
                    if key == "Indicators"
                    else field
                )
                val = display(actual_field, entry.get(field, ""))
                sheet.cell(row, col, val)
                lines = max(
                    lines,
                    len(
                        textwrap.wrap(
                            str(val),
                            width=max(
                                12,
                                int(
                                    sheet.column_dimensions[
                                        get_column_letter(col)
                                    ].width
                                )
                                - 3,
                            ),
                        )
                    ),
                )
            sheet.row_dimensions[row].height = max(36, lines * 16)
    reference = book.create_sheet(names["Reference"])
    reference.append([tr("Field"), tr("Available choice"), tr("Description")])
    for field, items in catalog.items():
        for value, display_value in items:
            description = (
                tr(C.output_title(value))
                if field == "output"
                else (
                    tr(C.key_activity_title(value))
                    if field == "key_activity_code"
                    else ""
                )
            )
            reference.append([S.label(field, tr), display_value, description])
    reference.column_dimensions["A"].width = 38
    reference.column_dimensions["B"].width = 65
    reference.column_dimensions["C"].width = 110
    reference.freeze_panes = "B2"
    reference.auto_filter.ref = reference.dimensions
    start["A34"] = tr(
        "Official classification titles are kept in their original language when no translation is available."
    )
    start.merge_cells("A34:D34")
    start["A34"].alignment = Alignment(wrap_text=True)
    start.row_dimensions[34].height = 30
    start["A35"] = tr("Reference lists")
    start["A35"].hyperlink = "#" + quote_sheetname(names["Reference"]) + "!A1"
    start["A35"].style = "Hyperlink"
    for sheet in book:
        sheet.sheet_view.rightToLeft = locale == "ar"
        sheet.sheet_view.showGridLines = False
        if sheet.sheet_state == "visible":
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.page_setup.orientation = "landscape"
            sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0
            sheet.print_options.horizontalCentered = True
            if sheet == start:
                sheet.print_area = "A1:D35"
                sheet.page_setup.fitToHeight = 1
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        if sheet not in (meta, opts):
            sheet.cell(
                2 if sheet.title in [names[k] for k in S.SHEETS] else 1, 1
            ).font = Font(name="Calibri", size=18, bold=True, color=blue)
    # Abrir una hoja útil: Actividades en el formato plano, la guía en v2.
    if flat:
        _format_flat(book, names, S, tr)
    book.active = book.index(book[names["Activities"]] if flat else start)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def parse_v2(book, schema=S):
    S = schema
    flat = getattr(S, "FLAT", False)
    from ckanext.theme_ejemplo import ihpix_bulk as B

    meta = book["_IHP_META"]
    if meta.max_row > 40 or meta.max_column > 2:
        raise ValueError("Workbook metadata is invalid")
    if meta["B1"].value != S.VERSION:
        raise ValueError("Unsupported template version; download the current template")
    locale = str(meta["B2"].value or "en")
    names = {
        str(row[0].value): str(row[1].value)
        for row in meta.iter_rows(min_row=5, max_row=4 + len(S.SHEET_LABELS), max_col=2)
        if row[0].value
    }
    if (
        set(names) != set(S.SHEET_LABELS)
        or len(set(names.values())) != len(names)
        or any(name not in book for name in names.values())
    ):
        raise ValueError("Worksheet structure changed; download the current template")
    if "_IHP_OPTIONS" not in book:
        raise ValueError("Workbook reference lists are missing")
    if set(book.sheetnames) != set(names.values()) | {"_IHP_META", "_IHP_OPTIONS"}:
        raise ValueError(
            "Unrecognized worksheet; move report data into the template sheets"
        )
    if book["_IHP_OPTIONS"].max_row > 5000 or book["_IHP_OPTIONS"].max_column > 3:
        raise ValueError("Workbook reference lists exceed template limits")
    options = {}
    for row in book["_IHP_OPTIONS"].iter_rows(min_row=2, max_col=3):
        field, display, value = [str(c.value or "") for c in row]
        if field:
            options.setdefault(field, {})[display.casefold()] = value
    errors = []
    sources = {}
    records = {name: [] for name in B.SHEETS}
    data_by_ref = {}

    def source(key, row, field):
        if flat and key == "Activities" and field == "row_ref":
            field = "title"
        col = list(S.SHEETS[key]).index(field) + 1
        return {
            "sheet": names[key],
            "row": row,
            "cell": get_column_letter(col) + str(row),
        }

    def error(key, row, field, message, ref=""):
        errors.append(
            dict(source(key, row, field), field=field, message=message, ref=ref)
        )

    def value(cell):
        val = cell.value
        if isinstance(val, (datetime.date, datetime.datetime)):
            return val.strftime("%Y-%m-%d")
        if isinstance(val, bool):
            return "yes" if val else "no"
        if isinstance(val, float) and val.is_integer():
            val = int(val)
        return "" if val is None else str(val).strip()

    def normalize(field, val):
        if not val:
            return ""
        category = choice_field(field)
        if field in F.MULTI_FORM_FIELDS or field == "types":
            return ";".join(
                normalize(
                    "country" if field == "member_states" else field + "_single",
                    part.strip(),
                )
                for part in val.split(";")
                if part.strip()
            )
        category = choice_field(field[:-7] if field.endswith("_single") else field)
        mapped = options.get(category, {}).get(val.casefold(), val)
        if category == "country":
            found = R.country_metadata(mapped)
            return found["country_title_en"] if found else mapped
        return mapped

    common = {}
    common_sources = {}
    for row, field in enumerate(S.COMMON, 7):
        if (
            meta.cell(row + 13, 1).value != field
            or book[names["Start"]].cell(row, 1).value != meta.cell(row + 13, 2).value
        ):
            raise ValueError(
                "Common reporting rows changed; keep the field labels in their original positions"
            )
        cell = book[names["Start"]].cell(row, 2)
        common[field] = normalize(field, value(cell))
        common_sources[field] = {
            "sheet": names["Start"],
            "row": row,
            "cell": "B" + str(row),
        }
        if cell.data_type == "f":
            errors.append(
                dict(
                    common_sources[field],
                    field=field,
                    message="Formulas are not allowed; paste values",
                    ref="",
                )
            )
    raw = {}
    for key, fields in S.SHEETS.items():
        sheet = book[names[key]]
        limit = B.MAX_ROWS * 8 if key == "Indicators" else B.MAX_ROWS
        if sheet.max_row > limit + 4 or sheet.max_column > len(fields):
            raise ValueError(
                "Worksheet dimensions exceed template limits: " + names[key]
            )
        if [c.value for c in next(sheet.iter_rows(max_row=1))] != list(fields):
            raise ValueError("Column structure changed: " + names[key])
        raw[key] = []
        for row, cells in enumerate(
            sheet.iter_rows(min_row=S.DATA_ROW, max_col=len(fields)), S.DATA_ROW
        ):
            entry = {
                field: normalize(field, value(cell))
                for field, cell in zip(fields, cells)
            }
            if not any(val for field, val in entry.items() if field != "row_ref"):
                continue
            if flat and key == "Activities":
                entry["row_ref"] = "A%03d" % (row - S.DATA_ROW + 1)
            for field, cell in zip(fields, cells):
                if cell.data_type == "f":
                    error(
                        key,
                        row,
                        field,
                        "Formulas are not allowed; paste values",
                        entry["row_ref"],
                    )
                if (
                    field in ("output", "key_activity_code")
                    and cell.value is not None
                    and cell.data_type == "n"
                ):
                    error(
                        key,
                        row,
                        field,
                        "Choose this code from the dropdown; codes must be text to preserve their digits.",
                        entry["row_ref"],
                    )
            raw[key].append((row, entry))
    for row, entry in raw["Activities"]:
        ref = entry["row_ref"]
        data = {field: "" for field in B.FIELDS}
        data.update(common)
        data.update(entry)
        if ref in data_by_ref or not ref:
            error(
                "Activities",
                row,
                "row_ref",
                "Use a unique, nonempty row reference",
                ref,
            )
            continue
        data_by_ref[ref] = (row, data)
        records["Activities"].append((row, data))
        for field in B.FIELDS:
            if field in common:
                sources["Activities", row, field] = common_sources[field]
            else:
                for owner, fields in S.SHEETS.items():
                    if field in fields and owner in (
                        "Activities",
                        "Institutions",
                        "Scope",
                    ):
                        sources["Activities", row, field] = (
                            source(owner, row, field)
                            if owner == "Activities"
                            else {
                                "sheet": names[owner],
                                "row": None,
                                "cell": "",
                                "missing_row": True,
                            }
                        )
                        break
        for code, mapping in S.INDICATORS.items() if not flat else ():
            for column, field in mapping.items():
                sources["Activities", row, field] = {
                    "sheet": names["Indicators"],
                    "row": None,
                    "cell": "",
                    "missing_row": True,
                    "indicator": code,
                }
    seen = set()
    for key in ("Institutions", "Scope", "Indicators"):
        for row, entry in raw.get(key, []):
            ref = entry["row_ref"]
            if ref not in data_by_ref:
                error(key, row, "row_ref", "No matching activity row", ref)
                continue
            main_row, data = data_by_ref[ref]
            indicator = entry.get("indicator", "")
            identity = (key, ref, indicator)
            if identity in seen:
                error(
                    key,
                    row,
                    "row_ref",
                    "Use only one row per activity in this section (one per indicator).",
                    ref,
                )
                continue
            seen.add(identity)
            if key == "Indicators":
                mapping = S.INDICATORS.get(indicator)
                if not mapping:
                    error(
                        key, row, "indicator", "Choose an indicator from the list", ref
                    )
                    continue
                for column, val in entry.items():
                    if column in ("row_ref", "indicator"):
                        continue
                    if column not in mapping:
                        if val:
                            error(
                                key,
                                row,
                                column,
                                "This field does not apply to the selected indicator",
                                ref,
                            )
                        continue
                    field = mapping[column]
                    data[field] = normalize(field, val)
                    sources["Activities", main_row, field] = source(key, row, column)
            else:
                if key == "Institutions" and (
                    entry["institution"] or entry["institution_org_id"]
                ):
                    data["institution"] = entry["institution"]
                    data["institution_org_id"] = entry["institution_org_id"]
                    for f in ("institution", "institution_org_id"):
                        sources["Activities", main_row, f] = source(key, row, f)
                for field, val in entry.items():
                    if field == "row_ref":
                        continue
                    if val or field not in common:
                        data[field] = val
                        sources["Activities", main_row, field] = source(key, row, field)
    if flat:
        titles = {}
        for ref, (_, data) in data_by_ref.items():
            titles.setdefault(data["title"], []).append(ref)
        for key in ("Contributions", "Supporters", "Partners", "Resources"):
            valid = []
            for row, entry in raw[key]:
                matches = titles.get(entry["row_ref"], [])
                if len(matches) != 1:
                    error(
                        key,
                        row,
                        "row_ref",
                        "Choose an activity title that appears exactly once in Activities",
                    )
                    continue
                entry["row_ref"] = matches[0]
                valid.append((row, entry))
            raw[key] = valid
    # La clasificación principal se inserta primero y nunca la sustituye un aporte secundario.
    for ref, (row, data) in data_by_ref.items():
        if data["output"]:
            records["Contributions"].append(
                (
                    row,
                    {
                        "row_ref": ref,
                        **{
                            field: data[field]
                            for field in (
                                "priority_area",
                                "output",
                                "key_activity_code",
                            )
                        },
                    },
                )
            )
            for field in B.SHEETS["Contributions"]:
                sources["Contributions", row, field] = source("Activities", row, field)
    # Los números virtuales evitan colisiones entre la clasificación principal y la hoja de aportes.
    for key in ("Contributions", "Supporters", "Partners", "Resources"):
        for row, entry in raw[key]:
            ref = entry["row_ref"]
            target = key
            virtual = row + B.MAX_ROWS + 10 if key == "Contributions" else row
            if key == "Contributions":
                relation = entry.pop("relation")
                if relation not in ("contribution", "synergy"):
                    error(
                        key,
                        row,
                        "relation",
                        "Choose Additional contribution or Synergy",
                        ref,
                    )
                    continue
                target = "Synergies" if relation == "synergy" else "Contributions"
                if target == "Contributions" and ref in data_by_ref:
                    main_row, data = data_by_ref[ref]
                    if not data["output"]:
                        error(
                            "Activities",
                            main_row,
                            "output",
                            "Complete the primary classification before adding contributions",
                            ref,
                        )
            elif key == "Supporters":
                selected = []
                for field in ("financial", "in_kind", "hosting"):
                    answer = entry.pop(field)
                    if answer and not F._gate_answered(answer):
                        error(key, row, field, "Please answer yes or no", ref)
                    elif C.normalize_bool(answer):
                        selected.append(field)
                entry["types"] = ";".join(selected)
            elif (
                key == "Resources"
                and not entry["target_kind"]
                and not entry["target_id"]
            ):
                entry["target_kind"] = "url"
            records[target].append((virtual, entry))
            for field in B.SHEETS[target]:
                origin = field if field in S.SHEETS[key] else "financial"
                sources[target, virtual, field] = source(key, row, origin)
    result = B.validate_records(
        records, sources=sources, version=S.VERSION, initial_errors=errors
    )
    if not result["rows"]:
        for issue in result["errors"]:
            if issue["message"] == "Add at least one activity":
                issue.update(source("Activities", 5, "title"))
                issue["field"] = "title"
    result.update(locale=locale, example=meta["B3"].value is True, common=common)
    return result


def _format_flat(book, names, schema, tr):
    """Guía breve y bloques horizontales; sin referencias que el usuario deba mantener."""
    from openpyxl.utils import get_column_letter

    main = book[names["Activities"]]
    book.move_sheet(main, offset=-book.index(main))
    for area in list(main.merged_cells.ranges):
        main.unmerge_cells(str(area))
    for row in (2, 3):
        for cell in main[row]:
            cell.value = None
    colors = ("005A96", "17606D", "335D82")
    column = 1
    for index, (title, fields) in enumerate(schema.GROUPS):
        end = column + len(fields) - 1
        main.merge_cells(start_row=2, start_column=column, end_row=2, end_column=end)
        cell = main.cell(2, column, tr(title))
        cell.fill = PatternFill("solid", fgColor=colors[index % len(colors)])
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=13)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        main.merge_cells(start_row=3, start_column=column, end_row=3, end_column=end)
        cell = main.cell(
            3,
            column,
            (
                tr(schema.GUIDE)
                if index == 0
                else tr(
                    "Leave unanswered fields blank. You can complete them in the draft."
                )
            ),
        )
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        for position in range(column, end + 1):
            main.cell(4, position).fill = PatternFill(
                "solid", fgColor=colors[index % len(colors)]
            )
        column = end + 1
    main.row_dimensions[2].height = 40
    main.row_dimensions[3].height = 36
    main.row_dimensions[4].height = 66
    main.freeze_panes = "B5"
    main.sheet_view.zoomScale = 85
    for col, field in enumerate(schema.SHEETS["Activities"], 1):
        if field in F.INT_FIELDS or field in F.BOOL_FIELDS:
            main.column_dimensions[get_column_letter(col)].width = 19
        elif field in ("start_date", "end_date"):
            main.column_dimensions[get_column_letter(col)].width = 18
    main.page_setup.fitToWidth = 0
    main.print_title_rows = "2:4"
    main.print_title_cols = "A:A"
    guide = book[names["Start"]]
    guide.freeze_panes = None
    for area in list(guide.merged_cells.ranges):
        guide.unmerge_cells(str(area))
    guide.delete_rows(1, guide.max_row)
    texts = [
        tr("Quick guide"),
        tr(
            "Fill in Activities: one activity per row, including its reporting details and indicators."
        ),
        tr(
            "Start with the activity title. Continue to the right and leave any unanswered fields blank."
        ),
        tr(
            "Use the dropdowns when available. No formulas or reference codes are needed."
        ),
        tr(
            "Repeat the contact and institution details for each activity. You can copy them from the previous row."
        ),
        tr(
            "The other sheets are optional. Use them only when an activity has additional contributions, support, partner organizations or resources."
        ),
        tr(
            "On an optional sheet, choose the activity title from the dropdown. Use distinct titles when adding these details."
        ),
        tr(
            "Upload the workbook, review the results and create drafts. Nothing is submitted automatically."
        ),
        tr(
            "The completed example is a separate file. Use the empty template for your real activities."
        ),
    ]
    for row, text in enumerate(texts, 1):
        guide.cell(row, 1, text)
        guide.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        guide.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="center")
        guide.row_dimensions[row].height = 42
    guide["A1"].font = Font(name="Calibri", size=18, bold=True, color="005A96")
    for row, (title, fields) in enumerate(schema.GROUPS, 11):
        column = schema.SHEETS["Activities"].index(fields[0]) + 1
        guide.cell(row, 1, tr(title))
        guide.cell(row, 1).hyperlink = (
            "#"
            + quote_sheetname(names["Activities"])
            + "!"
            + get_column_letter(column)
            + "5"
        )
        guide.cell(row, 1).style = "Hyperlink"
        guide.row_dimensions[row].height = 26
    guide.print_area = "A1:D" + str(guide.max_row)
